"""Reverse Engineering 101 — casual class deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Keep this in sync with the edited deck in reverse-engineering-101/.
# Do not put back lines Ajin deleted in PowerPoint.
OUTS = [
    Path(r"C:\Users\Public\Downloads\reverse-engineering-101\Reverse Engineering 101.pptx"),
    Path(r"C:\Users\Public\Downloads\Reverse Engineering 101.pptx"),
    Path(r"C:\Users\Public\Downloads\Dummy Reversing V2.pptx"),
    Path(r"C:\Users\Public\Downloads\Dummy Reversing V1.pptx"),
]
REG_IMG = Path(r"C:\Users\Public\Downloads\reverse-engineering-101\assets\x86_registers.png")

W, H = Inches(13.333), Inches(7.5)
BG = RGBColor(0x1A, 0x1A, 0x1A)
PANEL = RGBColor(0x24, 0x24, 0x24)
INK = RGBColor(0xF2, 0xF2, 0xF2)
DIM = RGBColor(0xB0, 0xB0, 0xB0)
ORANGE = RGBColor(0xE8, 0xA2, 0x3A)
GREEN = RGBColor(0x7C, 0xDA, 0x6A)
CODE_BG = RGBColor(0x10, 0x10, 0x10)
LINE = RGBColor(0x3A, 0x3A, 0x3A)


def run(p, text, size=18, color=INK, bold=False, font="Calibri"):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = font
    return r


def box(slide, l, t, w, h, fill=PANEL):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def txt(slide, l, t, w, h, text, size=18, color=INK, bold=False, font="Calibri", align=PP_ALIGN.LEFT):
    b = slide.shapes.add_textbox(l, t, w, h)
    tf = b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run(p, text, size, color, bold, font)
    return b


def lines(slide, l, t, w, h, items, size=18, color=INK, gap=10, font="Calibri"):
    b = slide.shapes.add_textbox(l, t, w, h)
    tf = b.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        if isinstance(item, str):
            run(p, item, size, color, False, font)
        else:
            text, c, bold, fnt, sz = item
            run(p, text, sz, c, bold, fnt)
    return b


def bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def footer(slide, n, total):
    box(slide, Inches(0), Inches(7.22), W, Inches(0.28), RGBColor(0x12, 0x12, 0x12))
    txt(slide, Inches(0.4), Inches(7.22), Inches(8), Inches(0.26),
        "Reverse Engineering 101   |   Ajin Deepak", 11, DIM)
    txt(slide, Inches(10.5), Inches(7.22), Inches(2.4), Inches(0.26),
        f"{n} / {total}", 11, DIM, align=PP_ALIGN.RIGHT)


def title_bar(slide, title):
    box(slide, Inches(0), Inches(0), Inches(0.12), H, ORANGE)
    txt(slide, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.55), title, 32, INK, True)


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    return s


def title_slide(s):
    box(s, Inches(0), Inches(0), Inches(0.14), H, ORANGE)
    txt(s, Inches(0.7), Inches(2.15), Inches(12), Inches(1.1),
        "Reverse Engineering 101", 46, INK, True)

    txt(s, Inches(0.7), Inches(4.5), Inches(11), Inches(0.3), "Presented by", 14, DIM)
    txt(s, Inches(0.7), Inches(4.85), Inches(11), Inches(0.45), "Ajin Deepak", 26, INK, True)
    txt(s, Inches(0.7), Inches(5.4), Inches(11), Inches(0.4),
        "AntiBot Engineer  @  NielsenIQ", 18, ORANGE)
    txt(s, Inches(0.7), Inches(6.55), Inches(12), Inches(0.45),
        "Parts of this came from an older deck and got updated.",
        14, DIM)


def syllabus_slide(s):
    title_bar(s, "What we're doing")
    items = [
        "What RE is, why people do it, where it's used",
        "Compile -> bytes -> disassembly",
        "Tools",
        "Hex edit, x32 assembly, GDB/GEF, then crackmes",
    ]
    for i, line in enumerate(items):
        y = 1.7 + i * 1.15
        txt(s, Inches(0.55), Inches(y), Inches(0.6), Inches(0.55), str(i + 1), 22, ORANGE, True)
        txt(s, Inches(1.2), Inches(y + 0.05), Inches(11), Inches(0.5), line, 24, INK)


def tools_slide(s):
    title_bar(s, "Tools")
    txt(s, Inches(0.5), Inches(1.0), Inches(12.3), Inches(0.55),
        "Today is 32-bit Linux ELF. Windows folks: WSL or a VM. Need gcc-multilib for -m32.", 18, DIM)
    tools = [
        ("GDB + GEF", "Debugger. GEF adds a clearer layout on top of GDB."),
        ("hexedit / Bless", "hexedit in the terminal, Bless if you want a GUI."),
        ("DIE", "Detect It Easy. Shows file type / packer / compiler."),
        ("Ghidra", "Free. Disassembly + decompiler."),
        ("IDA", "Paid. Same idea as Ghidra."),
        ("Binary Ninja", "Paid. Graph view is good if you have it."),
        ("JADX", "Android / Java. Not for today's ELF files."),
    ]
    for i, (name, why) in enumerate(tools):
        y = 1.58 + i * 0.74
        box(s, Inches(0.5), Inches(y), Inches(12.3), Inches(0.72), PANEL)
        txt(s, Inches(0.7), Inches(y + 0.16), Inches(2.6), Inches(0.42), name, 20, ORANGE, True)
        txt(s, Inches(3.4), Inches(y + 0.18), Inches(9.1), Inches(0.42), why, 18, INK)


def install_slide(s):
    title_bar(s, "Install GDB + GEF")
    txt(s, Inches(0.5), Inches(1.05), Inches(12.3), Inches(0.4),
        "Debian / Ubuntu / WSL. New terminal after GEF so it loads.", 17, DIM)
    box(s, Inches(0.5), Inches(1.55), Inches(12.3), Inches(5.4), PANEL)
    lines(s, Inches(0.75), Inches(1.75), Inches(11.8), Inches(5.0), [
        ("sudo apt update", GREEN, True, "Consolas", 20),
        ("sudo apt install -y build-essential gcc-multilib gdb python3 hexedit bless", GREEN, True, "Consolas", 18),
        ("bash -c \"$(curl -fsSL https://gef.blah.cat/sh)\"", GREEN, True, "Consolas", 20),
        ("", DIM, False, "Calibri", 10),
        ("Then open a new terminal and run:   gdb", DIM, False, "Calibri", 18),
        ("You should see a GEF banner. Same commands as vanilla GDB.", DIM, False, "Calibri", 18),
    ], gap=10)


def assemble_slide(s):
    title_bar(s, "Assemble / build the examples")
    txt(s, Inches(0.5), Inches(1.05), Inches(12.3), Inches(0.4),
        "32-bit. If gcc complains about -m32, you missed gcc-multilib.", 17, DIM)
    box(s, Inches(0.5), Inches(1.55), Inches(12.3), Inches(2.35), PANEL)
    txt(s, Inches(0.75), Inches(1.7), Inches(11.8), Inches(0.35), "all of them", 16, ORANGE, True)
    lines(s, Inches(0.75), Inches(2.15), Inches(11.8), Inches(1.5), [
        ("cd examples", GREEN, True, "Consolas", 20),
        ("make", GREEN, True, "Consolas", 20),
        ("file 01_mov", GREEN, True, "Consolas", 20),
    ], gap=6)
    box(s, Inches(0.5), Inches(4.1), Inches(12.3), Inches(2.8), PANEL)
    txt(s, Inches(0.75), Inches(4.25), Inches(11.8), Inches(0.35), "one file", 16, ORANGE, True)
    lines(s, Inches(0.75), Inches(4.7), Inches(11.8), Inches(2.0), [
        ("gcc -m32 -g -no-pie -o 01_mov 01_mov.s", GREEN, True, "Consolas", 20),
        ("gcc -m32 -g -no-pie -o 07_strings 07_strings.c", GREEN, True, "Consolas", 20),
        ("file should say ELF 32-bit", DIM, False, "Calibri", 16),
    ], gap=8)


def hexedit_slide(s):
    title_bar(s, "Hex-editing strings")
    txt(s, Inches(0.5), Inches(1.0), Inches(12.3), Inches(0.45),
        "Terminal: hexedit. GUI: Bless.  sudo apt install hexedit bless",
        17, DIM)
    box(s, Inches(0.5), Inches(1.55), Inches(7.9), Inches(5.4), PANEL)
    txt(s, Inches(0.75), Inches(1.7), Inches(7.4), Inches(0.35), "hexedit", 16, ORANGE, True)
    lines(s, Inches(0.75), Inches(2.15), Inches(7.4), Inches(4.6), [
        ("./07_strings", GREEN, True, "Consolas", 18),
        ("hexedit 07_strings", GREEN, True, "Consolas", 18),
        ("Ctrl+S     search  HELLO_FROM_RE101", GREEN, True, "Consolas", 18),
        ("Tab        hex / ASCII", GREEN, True, "Consolas", 18),
        ("type       BYE___FROM_RE101", GREEN, True, "Consolas", 18),
        ("F2         save", GREEN, True, "Consolas", 18),
        ("Ctrl+X     save and exit", GREEN, True, "Consolas", 18),
        ("bless 07_strings     GUI", GREEN, True, "Consolas", 18),
    ], gap=6)
    box(s, Inches(8.6), Inches(1.55), Inches(4.25), Inches(5.4), PANEL)
    txt(s, Inches(8.8), Inches(1.75), Inches(3.85), Inches(0.35), "program", 14, ORANGE, True)
    txt(s, Inches(8.8), Inches(2.2), Inches(3.85), Inches(0.7), "07_strings.c", 20, GREEN, True, "Consolas")
    txt(s, Inches(8.8), Inches(3.0), Inches(3.85), Inches(3.6),
        "HELLO_FROM_RE101 and BYE___FROM_RE101 are both 16 characters. Same length. Longer overwrites the next bytes (including the 0).",
        16, INK)


def what_is_slide(s):
    title_bar(s, "What is reverse engineering?")
    txt(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(1.1),
        "Looking at a finished program and figuring out how it works, without the original source.",
        22, INK)
    bullets = [
        "run it, see the output / input it asks for",
        "check strings and imports",
        "read the assembly, or the decompiler output",
        "debug it and watch the compare that decides pass/fail",
    ]
    for i, b in enumerate(bullets):
        txt(s, Inches(0.7), Inches(2.8 + i * 0.8), Inches(12), Inches(0.55), "-  " + b, 22, INK)


def why_slide(s):
    title_bar(s, "Why reverse engineering?")
    lines_ = [
        "You are lifeless, or you think you wanna be cool by cracking some crappy apps,",
        "or you wanna build some skills to get employed,",
        "or it's part of your job to upskill,",
        "or you are just interested.",
        "",
        "Idk which one.",
    ]
    for i, line in enumerate(lines_):
        y = 1.6 + i * 0.75
        color = DIM if not line else INK
        size = 18 if not line else 22
        txt(s, Inches(0.7), Inches(y), Inches(12), Inches(0.6), line, size, color)


def uses_slide(s):
    title_bar(s, "Where reverse engineering is used")
    rows = [
        ("Malware analysis", "You get a sample with no source. Goal is behavior: what it does, what it connects to, what it writes."),
        ("Appsec", "You are looking at a released binary. Find auth checks, parsers, and other risky spots."),
        ("Exploit development", "Something crashes or misbehaves. You need the real control flow, registers, and memory state."),
    ]
    for i, (title, body) in enumerate(rows):
        y = 1.35 + i * 1.8
        box(s, Inches(0.5), Inches(y), Inches(12.3), Inches(1.6), PANEL)
        txt(s, Inches(0.75), Inches(y + 0.2), Inches(11.8), Inches(0.4), title, 22, ORANGE, True)
        txt(s, Inches(0.75), Inches(y + 0.7), Inches(11.8), Inches(0.7), body, 18, INK)


def compile_slide(s):
    title_bar(s, "Compile -> bytes")
    txt(s, Inches(0.5), Inches(1.1), Inches(12.3), Inches(0.55),
        "C is for humans. The CPU runs bytes.", 20, DIM)
    steps = [
        ("1  Source", "int main() { return 1; }", "What you write."),
        ("2  Compiler", "gcc -m32  ->  ELF", "Builds the binary."),
        ("3  Bytes", "B8 01 00 00 00", "Machine code in .text."),
        ("4  CPU", "reads those bytes", "Decodes and runs them."),
    ]
    for i, (title, mid, why) in enumerate(steps):
        x = 0.45 + i * 3.2
        box(s, Inches(x), Inches(1.9), Inches(3.05), Inches(4.55), PANEL)
        txt(s, Inches(x + 0.18), Inches(2.1), Inches(2.7), Inches(0.7), title, 20, ORANGE, True)
        txt(s, Inches(x + 0.18), Inches(3.0), Inches(2.7), Inches(1.5), mid, 18, GREEN, True, "Consolas")
        txt(s, Inches(x + 0.18), Inches(4.7), Inches(2.7), Inches(1.4), why, 16, INK)


def tools_roles_slide(s):
    title_bar(s, "Compiler, disassembler, decompiler")
    txt(s, Inches(0.5), Inches(1.05), Inches(12.3), Inches(0.45),
        "These are not the same thing.", 18, DIM)
    cols = [
        ("Compiler", "gcc, clang", "Source -> binary.\n\nC/C++ becomes machine code + ELF headers.\n\nThis is how the file was made."),
        ("Disassembler", "objdump, Ghidra listing", "Bytes -> assembly.\n\nB8 01 00 00 00  ->  mov eax, 1\n\nMatches the bytes one-to-one."),
        ("Decompiler", "Ghidra / IDA / Binary Ninja", "Assembly -> C-like code.\n\nif (x == 1) ...\n\nHelpful, but it is a reconstruction. Names are usually gone."),
    ]
    for i, (title, who, body) in enumerate(cols):
        x = 0.45 + i * 4.25
        box(s, Inches(x), Inches(1.65), Inches(4.1), Inches(5.25), PANEL)
        txt(s, Inches(x + 0.2), Inches(1.85), Inches(3.7), Inches(0.45), title, 22, ORANGE, True)
        txt(s, Inches(x + 0.2), Inches(2.35), Inches(3.7), Inches(0.4), who, 15, GREEN, True)
        txt(s, Inches(x + 0.2), Inches(2.9), Inches(3.7), Inches(3.7), body, 16, INK)


def disasm_slide(s):
    title_bar(s, "How a disassembler reads those bytes")
    txt(s, Inches(0.5), Inches(1.05), Inches(12.3), Inches(0.5),
        "objdump / Ghidra / IDA / Binary Ninja all start by turning bytes into instructions.",
        18, DIM)
    box(s, Inches(0.5), Inches(1.65), Inches(12.3), Inches(1.7), PANEL)
    txt(s, Inches(0.75), Inches(1.8), Inches(11.8), Inches(0.4),
        "B8 01 00 00 00     ->     mov eax, 1", 26, GREEN, True, "Consolas")
    txt(s, Inches(0.75), Inches(2.45), Inches(11.8), Inches(0.7),
        "B8 = MOV EAX, imm32. Next 4 bytes = the number. Little endian, so 01 00 00 00 is 1. Total length: 5 bytes.",
        18, INK)
    bits = [
        ("Start address", "Starts at entry / a found function. Not at offset 0 of the file."),
        ("Variable length", "x86 instructions are 1 to 15 bytes. Opcode + ModR/M/SIB/imm set the size."),
        ("Wrong offset", "Start one byte off and the listing becomes junk."),
        ("Decompiler", "After the listing, tools try to rebuild C. Treat that as a helper, not ground truth."),
    ]
    for i, (h, d) in enumerate(bits):
        x = 0.5 + (i % 2) * 6.4
        y = 3.55 + (i // 2) * 1.65
        box(s, Inches(x), Inches(y), Inches(6.2), Inches(1.5), PANEL)
        txt(s, Inches(x + 0.2), Inches(y + 0.12), Inches(5.8), Inches(0.35), h, 18, ORANGE, True)
        txt(s, Inches(x + 0.2), Inches(y + 0.55), Inches(5.8), Inches(0.8), d, 15, INK)


def asm_slide(s):
    title_bar(s, "Assembly")
    txt(s, Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.7),
        "Each instruction does one small thing: move, add, compare, jump.",
        20, INK)
    box(s, Inches(0.5), Inches(2.2), Inches(6.0), Inches(4.3), PANEL)
    txt(s, Inches(0.75), Inches(2.4), Inches(5.5), Inches(0.4), "Syntax we use", 16, ORANGE, True)
    lines(s, Inches(0.75), Inches(3.0), Inches(5.5), Inches(3.2), [
        ("mov eax, 1", GREEN, True, "Consolas", 22),
        ("destination on the left", DIM, False, "Calibri", 18),
        ("", DIM, False, "Calibri", 10),
        ("Our .s files use this form.", DIM, False, "Calibri", 18),
        ("Same encoding as in the binary.", DIM, False, "Calibri", 18),
    ], gap=6)
    box(s, Inches(6.8), Inches(2.2), Inches(6.0), Inches(4.3), PANEL)
    txt(s, Inches(7.05), Inches(2.4), Inches(5.5), Inches(0.4), "32-bit", 16, ORANGE, True)
    lines(s, Inches(7.05), Inches(3.0), Inches(5.5), Inches(3.2), [
        "EAX, EBX, ECX, EDX, ESI, EDI.",
        "ESP = stack pointer.",
        "EBP = frame pointer.",
        "EIP = next instruction.",
        "gcc -m32  for the examples.",
    ], size=18, gap=10)


def registers_slide(s):
    title_bar(s, "Registers")
    txt(s, Inches(0.5), Inches(1.0), Inches(12.3), Inches(0.4),
        "Fast storage inside the CPU. Most work happens here.", 18, DIM)
    if REG_IMG.exists():
        s.shapes.add_picture(str(REG_IMG), Inches(0.45), Inches(1.5), Inches(7.3), Inches(5.4))
    notes = [
        ("EAX", "return value, math"),
        ("EBX", "general purpose"),
        ("ECX", "loop counter"),
        ("EDX", "extra math / IO"),
        ("ESI EDI", "source / dest for copies"),
        ("ESP", "stack pointer"),
        ("EBP", "frame pointer"),
        ("EIP", "next instruction"),
    ]
    for i, (n, d) in enumerate(notes):
        y = 1.5 + i * 0.65
        txt(s, Inches(8.0), Inches(y), Inches(2.2), Inches(0.5), n, 15, ORANGE, True, "Consolas")
        txt(s, Inches(10.2), Inches(y), Inches(2.7), Inches(0.5), d, 15, INK)


def debug_slide(s):
    title_bar(s, "Debugging")
    txt(s, Inches(0.5), Inches(1.1), Inches(12.3), Inches(0.7),
        "Ghidra is static. Debugging means running the program and stopping it where you care.",
        20, INK)
    left = [
        ("Breakpoints", "stop at main, or on a compare"),
        ("Step", "one instruction at a time"),
        ("Inspect", "registers, stack, pointers / strings"),
        ("Continue", "run until the next break"),
    ]
    for i, (h, d) in enumerate(left):
        y = 2.0 + i * 1.15
        box(s, Inches(0.5), Inches(y), Inches(12.3), Inches(1.02), PANEL)
        txt(s, Inches(0.75), Inches(y + 0.12), Inches(3.2), Inches(0.7), h, 22, ORANGE, True)
        txt(s, Inches(4.2), Inches(y + 0.25), Inches(8.3), Inches(0.55), d, 20, INK)


def gdb_slide(s):
    title_bar(s, "GDB commands")
    txt(s, Inches(0.5), Inches(1.0), Inches(12.3), Inches(0.4),
        "Practice on 01_mov first.", 17, DIM)
    cmds = [
        ("gdb ./01_mov", "open binary"),
        ("break main", "breakpoint at main"),
        ("run", "start"),
        ("disassemble", "show code here"),
        ("stepi    (si)", "step into"),
        ("nexti    (ni)", "step over"),
        ("info registers", "show registers"),
        ("print /x $eax", "print eax in hex"),
        ("x/10i $eip", "next 10 instructions"),
        ("x/8xw $esp", "stack dump"),
        ("x/s $eax", "string at eax"),
        ("continue    (c)", "continue"),
        ("quit", "exit gdb"),
    ]
    for i, (cmd, why) in enumerate(cmds):
        col = i // 7
        row = i % 7
        x = 0.45 + col * 6.45
        y = 1.5 + row * 0.80
        box(s, Inches(x), Inches(y), Inches(6.25), Inches(0.64), PANEL)
        txt(s, Inches(x + 0.15), Inches(y + 0.12), Inches(3.7), Inches(0.4), cmd, 14, GREEN, True, "Consolas")
        txt(s, Inches(x + 3.9), Inches(y + 0.12), Inches(2.2), Inches(0.4), why, 14, DIM)


def gdb_flow_slide(s):
    title_bar(s, "GDB walkthrough")
    steps = [
        "cd examples && make",
        "gdb ./01_mov",
        "break main",
        "run",
        "disassemble",
        "si   then   info registers",
    ]
    for i, line in enumerate(steps):
        y = 1.25 + i * 0.9
        box(s, Inches(0.5), Inches(y), Inches(12.3), Inches(0.78), PANEL)
        txt(s, Inches(0.7), Inches(y + 0.16), Inches(0.7), Inches(0.5), str(i + 1), 22, ORANGE, True)
        txt(s, Inches(1.5), Inches(y + 0.18), Inches(11), Inches(0.5), line, 20, INK, font="Consolas" if i else "Calibri")


def instr_slide(s, title, blurb, demo_file, rows, chal_file, challenge, gdb_hint):
    title_bar(s, title)
    txt(s, Inches(0.5), Inches(0.95), Inches(12.3), Inches(0.32), blurb, 14, DIM)

    box(s, Inches(0.5), Inches(1.3), Inches(6.5), Inches(3.5), PANEL)
    txt(s, Inches(0.65), Inches(1.4), Inches(6.1), Inches(0.28), "example  (" + demo_file + ")", 12, ORANGE, True)
    txt(s, Inches(0.65), Inches(1.7), Inches(3.0), Inches(0.25), "asm", 11, DIM, True)
    txt(s, Inches(3.7), Inches(1.7), Inches(3.0), Inches(0.25), "closest C", 11, DIM, True)
    for i, (cmd, c) in enumerate(rows):
        y = 2.05 + i * 0.62
        txt(s, Inches(0.65), Inches(y), Inches(3.0), Inches(0.45), cmd, 15, GREEN, True, "Consolas")
        txt(s, Inches(3.7), Inches(y), Inches(3.1), Inches(0.45), c, 14, INK, False, "Consolas")

    box(s, Inches(7.2), Inches(1.3), Inches(5.65), Inches(3.5), PANEL)
    txt(s, Inches(7.4), Inches(1.4), Inches(5.25), Inches(0.28), "challenge  (" + chal_file + ")", 12, ORANGE, True)
    txt(s, Inches(7.4), Inches(1.8), Inches(5.25), Inches(2.7), challenge, 15, INK)

    box(s, Inches(0.5), Inches(5.0), Inches(12.35), Inches(1.85), PANEL)
    txt(s, Inches(0.7), Inches(5.1), Inches(12), Inches(0.28), "gdb on the challenge binary (not the example)", 13, ORANGE, True)
    txt(s, Inches(0.7), Inches(5.45), Inches(12), Inches(1.25), gdb_hint, 15, GREEN, True, "Consolas")


def mov_slide(s):
    instr_slide(
        s, "mov",
        "Like assignment in C. Source stays the same. [ ] means memory.",
        "01_mov",
        [
            ("mov eax, 1", "eax = 1;"),
            ("mov eax, ebx", "eax = ebx;"),
            ("mov eax, [ebx]", "eax = *ebx;"),
            ("mov [ebx], eax", "*ebx = eax;"),
        ],
        "ch_mov",
        "Different numbers than the example.\n\nAt label check:\n  eax = ?\n  ecx = ?\n\nUse gdb. Do not guess from this slide.",
        "make challenges\ngdb ./ch_mov\nbreak check\nrun\nprint /x $eax\nprint /x $ecx",
    )


def arith_slide(s):
    instr_slide(
        s, "add  sub  inc  dec",
        "Same idea as += -= ++ --. Flags get updated. Note: inc/dec do not change CF.",
        "02_arith",
        [
            ("add eax, 5", "eax += 5;"),
            ("sub eax, 3", "eax -= 3;"),
            ("inc eax", "eax++;"),
            ("dec eax", "eax--;"),
        ],
        "ch_arith",
        "Different sequence than the example.\n\nAt label check:\n  eax = ?",
        "gdb ./ch_arith\nbreak check\nrun\nprint $eax",
    )


def xor_slide(s):
    instr_slide(
        s, "xor  +  zero flag",
        "Bitwise XOR. xor reg, reg clears that register.",
        "03_xor",
        [
            ("xor eax, eax", "eax = 0;"),
            ("xor eax, 0x55", "eax ^= 0x55;"),
            ("xor ebx, ebx", "ebx = 0;"),
        ],
        "ch_xor",
        "Different values than the example.\n\nAt label check:\n  eax = ?\n  is ZF set?\n\nZF: print ($eflags & 0x40)\n(non-zero means set)",
        "gdb ./ch_xor\nbreak check\nrun\nprint /x $eax\nprint ($eflags & 0x40)",
    )


def jump_slide(s):
    instr_slide(
        s, "cmp   jmp   je   jne",
        "cmp sets flags only. je/jne look at ZF.",
        "04_cmp_jmp",
        [
            ("cmp eax, ebx", "/* ZF if equal */"),
            ("jmp label", "goto label;"),
            ("je  /  jz", "if equal goto;"),
            ("jne / jnz", "if not equal goto;"),
        ],
        "ch_jmp",
        "Different values / branches.\n\nAt label check:\n  eax = ?\n  which path ran?",
        "gdb ./ch_jmp\nbreak check\nrun\nprint $eax",
    )


def stack_slide(s):
    instr_slide(
        s, "push  /  pop",
        "No direct C operator. Used for args and saved values.",
        "05_stack",
        [
            ("push eax", "/* push onto stack */"),
            ("pop ecx", "/* pop into ecx */"),
            ("x/2xw $esp", "/* dump stack */"),
        ],
        "ch_stack",
        "Different constants than the example.\n\nAt before_pop:\n  [esp] = ?\n  [esp+4] = ?\n\nAt check:\n  ecx = ?\n  edx = ?",
        "gdb ./ch_stack\nbreak before_pop\nrun\nx/2xw $esp\nbreak check\ncontinue\nprint /x $ecx\nprint /x $edx",
    )


def call_slide(s):
    instr_slide(
        s, "call  /  ret",
        "call jumps to a function. Return value usually in eax.",
        "06_call",
        [
            ("push 41", "/* argument */"),
            ("call add_one", "eax = add_one(41);"),
            ("add esp, 4", "/* cdecl cleanup */"),
            ("ret", "return;"),
        ],
        "ch_call",
        "Different function / argument.\n\nAt label check\n(right after call returns):\n  eax = ?",
        "gdb ./ch_call\nbreak check\nrun\nprint $eax",
    )


def ghidra_slide(s):
    title_bar(s, "Ghidra")
    txt(s, Inches(0.5), Inches(1.0), Inches(12.3), Inches(0.4),
        "Open 08_ghidra_flow first, then the crackme.", 18, DIM)
    bits = [
        ("New project", "File -> New Project, then Import File on 08_ghidra_flow."),
        ("Analyze", "Say Yes to analysis. Defaults are fine."),
        ("Two views", "Listing = asm. Decompile = reconstructed C."),
        ("Find main", "Symbol Tree -> Functions -> main / add_one / is_ok."),
        ("Then GDB", "Copy the address or name and break on it."),
    ]
    for i, (h, d) in enumerate(bits):
        y = 1.5 + i * 1.05
        box(s, Inches(0.5), Inches(y), Inches(12.3), Inches(0.95), PANEL)
        txt(s, Inches(0.7), Inches(y + 0.22), Inches(2.6), Inches(0.5), h, 18, ORANGE, True)
        txt(s, Inches(3.5), Inches(y + 0.22), Inches(9.0), Inches(0.55), d, 18, INK)


def crackme1(s):
    title_bar(s, "Crackme 1")
    txt(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(0.7),
        "Live. Linux binary.", 22, INK)
    bits = [
        "Run it and note what it asks for.",
        "strings ./crackme1",
        "gdb: break main, run, find the compare",
        "No patching. Explain what it checks.",
    ]
    for i, line in enumerate(bits):
        y = 2.15 + i * 1.05
        box(s, Inches(0.5), Inches(y), Inches(12.3), Inches(0.9), PANEL)
        txt(s, Inches(0.8), Inches(y + 0.22), Inches(11.7), Inches(0.5), line, 22, INK)


def crackme2(s):
    title_bar(s, "Crackme 2")
    txt(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(0.7),
        "Same approach. A bit more logic.", 22, INK)
    bits = [
        "run -> strings -> break -> watch cmp",
        "Expect more than one check.",
        "If stuck, open it in Ghidra/IDA, find the function, break there.",
        "No patching. Walk through the path.",
    ]
    for i, line in enumerate(bits):
        y = 2.15 + i * 1.05
        box(s, Inches(0.5), Inches(y), Inches(12.3), Inches(0.9), PANEL)
        txt(s, Inches(0.8), Inches(y + 0.22), Inches(11.7), Inches(0.5), line, 22, INK)


def end_slide(s):
    box(s, Inches(0), Inches(0), Inches(0.14), H, ORANGE)
    txt(s, Inches(0.7), Inches(2.1), Inches(12), Inches(0.8), "Done.", 36, INK, True)
    txt(s, Inches(0.7), Inches(3.1), Inches(12), Inches(0.6),
        "Rebuild the examples, step them again, then redo both crackmes without notes.",
        20, DIM)
    txt(s, Inches(0.7), Inches(4.4), Inches(12), Inches(0.4), "Ajin Deepak", 22, INK, True)
    txt(s, Inches(0.7), Inches(4.9), Inches(12), Inches(0.4),
        "AntiBot Engineer  @  NielsenIQ", 18, ORANGE)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    # Order matches the edited deck (tools_roles before compile->bytes).
    # why_slide sits right after what_is_slide.
    fns = [
        title_slide, syllabus_slide, what_is_slide, why_slide, uses_slide,
        tools_roles_slide, compile_slide, disasm_slide,
        tools_slide, install_slide, assemble_slide, hexedit_slide, asm_slide,
        registers_slide, debug_slide, gdb_slide, gdb_flow_slide, mov_slide,
        arith_slide, xor_slide, jump_slide, stack_slide, call_slide,
        ghidra_slide, crackme1, crackme2, end_slide,
    ]
    total = len(fns)
    for i, fn in enumerate(fns, 1):
        s = new_slide(prs)
        fn(s)
        if fn is not title_slide and fn is not end_slide:
            footer(s, i, total)
    saved = []
    for out in OUTS:
        try:
            prs.save(out)
            saved.append(str(out))
        except PermissionError:
            alt = out.with_name(out.stem + " - updated.pptx")
            prs.save(alt)
            saved.append(f"{out} locked, wrote {alt}")
    print("saved:")
    for p in saved:
        print(" ", p)


if __name__ == "__main__":
    build()
